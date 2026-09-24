"""DOC/DOCX adapter for LLM-ready content."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


class DocParserAdapter:
    """Adapter for Word-like files."""

    def supports(self, content_type: str, filename: str) -> bool:
        lower = filename.lower()
        return lower.endswith((".doc", ".docx")) or content_type in {
            "application/msword",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        }

    def adapt(self, file_bytes: bytes, filename: str, content_type: str) -> dict[str, object]:
        # Let parse errors propagate (do NOT fall back to utf-8 decode, which produced
        # garbage for binary .doc files); extract_content reports parse_ok=False instead.
        doc = Document(BytesIO(file_bytes))

        lines: list[str] = []
        tables: list[list[list[str]]] = []
        for block in self._iter_block_items(doc):
            if isinstance(block, Paragraph):
                if block.text.strip():
                    lines.append(block.text)
                continue
            # Many resume/report templates lay their entire body out inside a table
            # (common two-column layout), not paragraphs — reading only .paragraphs
            # silently returned empty text for those documents.
            rows = [[cell.text for cell in row.cells] for row in block.rows]
            tables.append(rows)
            for row in rows:
                row_text = "\t".join(cell for cell in row if cell.strip())
                if row_text:
                    lines.append(row_text)

        return {
            "adapter": "doc",
            "filename": filename,
            "content_type": content_type,
            "text": "\n".join(lines),
            "images": [],
            "tables": tables,
        }

    @staticmethod
    def _iter_block_items(doc: Document):
        """Yield each top-level Paragraph/Table in document order.

        ``doc.paragraphs``/``doc.tables`` split the two apart and lose their
        relative order; walking the body's XML children directly preserves the
        same reading order the document actually displays.
        """
        for child in doc.element.body.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, doc)
            elif child.tag == qn("w:tbl"):
                yield Table(child, doc)
